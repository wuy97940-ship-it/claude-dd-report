# -*- coding: utf-8 -*-
# dd_report_builder.py <content.json> <template.docx> <output.docx>
# 尽调报告生成器 v3.2（格式母版法；v3.2新增title/sub块用于无封面简版说明）
# 原理：以公司已定稿的尽调报告为格式母版，克隆其段落格式样本(pPr)，
#       清空正文后按content.json重建，100%继承母版字体/缩进/行距/表格样式。
# content.json 格式：
#   {"meta": {"cover_title1": "...", "cover_title2": "...", "cover_dept": "...", "cover_date": "..."},
#    "blocks": [ {"t":"cover"|"toc"|"h1"|"h2"|"br"|"p"|"pni"|"note"|"unit"|"tbl"|"pb"|"title"|"sub", ...} ]}
#   title块(v3.2): 文档大标题(居中加粗,默认16pt,size可覆盖)——用于无封面的简版说明;
#   sub块(v3.2):   居中副行(默认10.5pt不加粗),如编制日期行。
#   tbl块: {"t":"tbl","fs":9,"cols":[...],"rows":[[...]]}
import sys, json, copy, re
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH

EASTASIA = '宋体'
ASCII_FONT = 'Times New Roman'
CIRCLED_RE = re.compile('[①-⑳㉑-㉟㊱-㊿]')   # 圈号数字①-⑳等，按终审人偏好须用宋体渲染

def classify_cell(text):
    """用户对齐规则：数字靠右，文字靠左，符号/日期/序号居中"""
    s = text.strip()
    if not s:
        return WD_ALIGN_PARAGRAPH.CENTER
    if s in ('合计', '小计'):
        return WD_ALIGN_PARAGRAPH.CENTER
    if re.fullmatch(r'\d{1,3}', s):                       # 序号
        return WD_ALIGN_PARAGRAPH.CENTER
    if re.fullmatch(r'\d{4}[-/]\d{1,2}[-/]\d{1,2}', s):   # 日期
        return WD_ALIGN_PARAGRAPH.CENTER
    if re.fullmatch(r'[—\-/、%]+', s):                    # 纯符号
        return WD_ALIGN_PARAGRAPH.CENTER
    if re.fullmatch(r'-?[\d,]+(\.\d+)?%?', s):            # 数字
        return WD_ALIGN_PARAGRAPH.RIGHT
    return WD_ALIGN_PARAGRAPH.LEFT                        # 文字

def _styled_run(p, text, size, bold, hansi, hint=False):
    r = p.add_run(text)
    rPr = r._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:ascii'), ASCII_FONT)
    rFonts.set(qn('w:hAnsi'), hansi)
    rFonts.set(qn('w:eastAsia'), EASTASIA)
    if hint:
        # 圈码①-⑳是字体归属歧义字符：必须带 w:hint="eastAsia"，否则被西文字体接管、显示偏大
        # （selfcheck[1]门禁项；20261002某正极材料客户壬轮首次在builder侧根治）
        rFonts.set(qn('w:hint'), 'eastAsia')
    if size is not None:
        r.font.size = Pt(size)
    if bold is not None:
        r.font.bold = bold
    return r

def styled_run(p, text, size=None, bold=None):
    """添加run；圈号数字①-⑳单独成run并设hAnsi=宋体（终审人偏好：序号用宋体好看），
    其余西文/数字用Times New Roman、中文用宋体。"""
    if CIRCLED_RE.search(text):
        last = 0
        for m in CIRCLED_RE.finditer(text):
            if m.start() > last:
                _styled_run(p, text[last:m.start()], size, bold, ASCII_FONT)
            _styled_run(p, m.group(), size, bold, EASTASIA, hint=True)
            last = m.end()
        if last < len(text):
            _styled_run(p, text[last:], size, bold, ASCII_FONT)
    else:
        _styled_run(p, text, size, bold, ASCII_FONT)
    return None

def strip_outline(ppr):
    for ol in ppr.findall(qn('w:outlineLvl')):
        ppr.remove(ol)

def new_para(doc, ppr_xml=None):
    p = doc.add_paragraph()
    if ppr_xml is not None:
        old = p._element.find(qn('w:pPr'))
        if old is not None:
            p._element.remove(old)
        p._element.insert(0, copy.deepcopy(ppr_xml))
    return p

def add_outline(p, level):
    pPr = p._element.find(qn('w:pPr'))
    if pPr is None:
        pPr = OxmlElement('w:pPr')
        p._element.insert(0, pPr)
    ol = OxmlElement('w:outlineLvl')
    ol.set(qn('w:val'), str(level))
    pPr.append(ol)

def make_ppr(align=None, indent_chars=None, line_15=True):
    """合成兜底pPr（母版缺样本时使用）"""
    pPr = OxmlElement('w:pPr')
    if align:
        jc = OxmlElement('w:jc'); jc.set(qn('w:val'), align); pPr.append(jc)
    if indent_chars:
        ind = OxmlElement('w:ind')
        ind.set(qn('w:firstLine'), str(indent_chars * 240))
        ind.set(qn('w:firstLineChars'), str(indent_chars * 100))
        pPr.append(ind)
    if line_15:
        sp = OxmlElement('w:spacing')
        sp.set(qn('w:line'), '360'); sp.set(qn('w:lineRule'), 'auto')
        pPr.append(sp)
    return pPr

def add_toc_field(doc, ppr_center):
    t = new_para(doc, ppr_center)
    styled_run(t, '目录', size=16, bold=True)
    p = doc.add_paragraph()
    r1 = p.add_run()
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin'); r1._element.append(f1)
    r2 = p.add_run()
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve')
    it.text = 'TOC \\o "1-2" \\h \\z \\u'
    r2._element.append(it)
    r3 = p.add_run()
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'separate'); r3._element.append(f2)
    styled_run(p, '（打开Word后：右键此处→更新域，即可自动生成目录及页码）')
    r5 = p.add_run()
    f3 = OxmlElement('w:fldChar'); f3.set(qn('w:fldCharType'), 'end'); r5._element.append(f3)
    br = doc.add_paragraph()
    br.add_run().add_break(WD_BREAK.PAGE)

def add_table(doc, block):
    cols = block['cols']
    rows = block['rows']
    fs = block.get('fs', 9)
    # 表格标题：居中加粗
    if block.get('title'):
        tp = doc.add_paragraph()
        tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        styled_run(tp, block['title'], size=12, bold=True)
    tb = doc.add_table(rows=len(rows) + 1, cols=len(cols))
    for style_name in ('Table Grid', '网格型'):   # v3.2: 中文模板样式名fallback
        try:
            tb.style = doc.styles[style_name]
            break
        except KeyError:
            continue
    tb.autofit = True
    # 三线表边框：顶线medium(1.5pt)+底线thick(2.25pt)大黑粗线，内部hair细线，无左右边线
    tblPr = tb._element.tblPr
    # 表格按窗口自动调整宽度（终审人指令：所有表格统一按窗口自动调整）
    tblW = OxmlElement('w:tblW'); tblW.set(qn('w:w'), '0'); tblW.set(qn('w:type'), 'auto')
    tblPr.append(tblW)
    layout = OxmlElement('w:tblLayout'); layout.set(qn('w:type'), 'autofit')
    tblPr.append(layout)
    jc = OxmlElement('w:jc'); jc.set(qn('w:val'), 'center')   # 表格水平居中
    tblPr.append(jc)
    borders = OxmlElement('w:tblBorders')
    for tag, sz in [('top', '12'), ('bottom', '18'), ('insideH', '4'), ('insideV', '4')]:
        el = OxmlElement('w:' + tag)
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), sz); el.set(qn('w:color'), '000000')
        borders.append(el)
    for tag in ('left', 'right'):
        el = OxmlElement('w:' + tag)
        el.set(qn('w:val'), 'none'); el.set(qn('w:sz'), '0'); el.set(qn('w:space'), '0')
        borders.append(el)
    tblPr.append(borders)
    # 表头：跨页重复、居中加粗
    trPr = tb.rows[0]._element.get_or_add_trPr()
    th = OxmlElement('w:tblHeader')
    th.set(qn('w:val'), 'true')
    trPr.append(th)
    for j, h in enumerate(cols):
        p = tb.rows[0].cells[j].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        styled_run(p, h, size=fs, bold=True)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            if j >= len(cols):
                break
            cell = tb.rows[i].cells[j]
            parts = str(val).split('\n')
            p0 = cell.paragraphs[0]
            p0.alignment = classify_cell(parts[0])
            styled_run(p0, parts[0], size=fs, bold=False)
            for extra in parts[1:]:
                px = cell.add_paragraph()
                px.alignment = classify_cell(extra)
                styled_run(px, extra, size=fs, bold=False)
    for r in tb.rows:
        for c in r.cells:
            tcPr = c._element.get_or_add_tcPr()
            va = OxmlElement('w:vAlign')
            va.set(qn('w:val'), 'center')
            tcPr.append(va)
    return tb

def _is_toc_cache(pPr, text):
    """识别母版TOC域缓存条目段落（带点线制表位且行尾为页码）。
    中文Word样式名为"目录 2"等，用style_name判断' toc'会漏（20260924实测），
    必须用结构特征排除——否则br/p等模板抓到缓存条目，正文字段落被挂
    目录样式+左缩进480（客户丙之子公司丙1初稿踩坑，终审人批注抓出）。"""
    if pPr is None:
        return False
    tabs = pPr.findall(qn('w:tabs') + '/' + qn('w:tab'))
    has_dot = any(tb.get(qn('w:leader')) == 'dot' for tb in tabs)
    import re as _re
    return has_dot and bool(_re.search(r'\d{1,4}\s*$', text or ''))


def capture_exemplars(doc):
    """从格式母版自动识别格式样本，无需硬编码任何公司名"""
    exp = {}
    for idx, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        pPr = p._element.find(qn('w:pPr'))
        if pPr is None:
            continue
        if _is_toc_cache(pPr, t):
            continue  # 跳过TOC缓存条目，模板只能取正文段落
        fmt = p.paragraph_format
        style_name = p.style.name if p.style is not None else ''
        if 'blank' not in exp and not t:
            exp['blank'] = copy.deepcopy(pPr)
            continue
        if not t:
            continue
        if ('cover' not in exp and fmt.alignment == WD_ALIGN_PARAGRAPH.CENTER
                and any(r.font.bold for r in p.runs)):
            exp['cover'] = copy.deepcopy(pPr)
        if 'h1' not in exp and style_name in ('标题1', '标题 1', 'Heading 1'):
            exp['h1'] = copy.deepcopy(pPr)
        if 'h2' not in exp and (t.startswith('（一）') or t.startswith('(一)')) and 'toc' not in style_name.lower():
            exp['h2'] = copy.deepcopy(pPr)
        if 'br' not in exp and t.startswith('【'):
            exp['br'] = copy.deepcopy(pPr)
        if 'note' not in exp and t.startswith('备注'):
            exp['note'] = copy.deepcopy(pPr)
        if 'unit' not in exp and t.startswith('单位'):
            exp['unit'] = copy.deepcopy(pPr)
        if ('p' not in exp and len(t) > 30 and not t.startswith(('（', '(', '【', '备注', '单位'))
                and fmt.first_line_indent is not None and fmt.first_line_indent >= Pt(20)
                and fmt.alignment != WD_ALIGN_PARAGRAPH.CENTER
                and 'toc' not in style_name.lower()):
            exp['p'] = copy.deepcopy(pPr)
    # 兜底合成
    fallbacks = {
        'cover': lambda: make_ppr(align='center'),
        'blank': lambda: make_ppr(),
        'h1': lambda: make_ppr(),
        'h2': lambda: make_ppr(align='both', indent_chars=2),
        'br': lambda: make_ppr(align='both'),
        'p': lambda: make_ppr(align='both', indent_chars=2),
        'note': lambda: make_ppr(align='both'),
        'unit': lambda: make_ppr(align='right', indent_chars=2),
    }
    for k, fn in fallbacks.items():
        if k not in exp:
            exp[k] = fn()
            print('INFO exemplar synthesized: %s' % k)
    for k in exp:
        strip_outline(exp[k])
    return exp

def main():
    content_path, template_path, out_path = sys.argv[1], sys.argv[2], sys.argv[3]
    with open(content_path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    doc = Document(template_path)
    exp = capture_exemplars(doc)

    body = doc.element.body
    for child in list(body):
        if child.tag != qn('w:sectPr'):
            body.remove(child)

    meta = data.get('meta', {})

    def cover_line(text):
        p = new_para(doc, exp['cover'])
        styled_run(p, text, size=18, bold=True)

    for block in data['blocks']:
        t = block['t']
        if t == 'cover':
            new_para(doc, exp['blank']); new_para(doc, exp['blank'])
            cover_line(meta.get('cover_title1', ''))
            cover_line(meta.get('cover_title2', ''))
            for _ in range(6):
                new_para(doc, exp['blank'])
            cover_line(meta.get('cover_dept', ''))
            new_para(doc, exp['blank'])
            cover_line(meta.get('cover_date', ''))
            br = doc.add_paragraph()
            br.add_run().add_break(WD_BREAK.PAGE)
        elif t == 'toc':
            add_toc_field(doc, exp['cover'])
        elif t == 'h1':
            p = new_para(doc, exp['h1'])
            styled_run(p, block['x'])
            add_outline(p, 0)
        elif t == 'h2':
            p = new_para(doc, exp['h2'])
            styled_run(p, block['x'], bold=True)
            add_outline(p, 1)
        elif t == 'br':
            p = new_para(doc, exp['br'])
            styled_run(p, block['x'], bold=True)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif t == 'p':
            p = new_para(doc, exp['p'])
            styled_run(p, block['x'])
        elif t == 'pni':
            p = new_para(doc, exp['br'])
            styled_run(p, block['x'])
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif t == 'note':
            p = new_para(doc, exp['note'])
            styled_run(p, block['x'], size=10.5)
        elif t == 'unit':
            p = new_para(doc, exp['unit'])
            styled_run(p, block['x'])
        elif t == 'tbl':
            add_table(doc, block)
        elif t == 'img':
            # 行业趋势图表标配(v5.6)：图片居中,默认宽13.5cm,title为图注(备注样式五号字)
            from docx.shared import Cm as _Cm
            _p = doc.add_paragraph()
            _p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _run = _p.add_run()
            _run.add_picture(block['path'], width=_Cm(block.get('width_cm', 13.5)))
            if block.get('title'):
                _cap = new_para(doc, exp['note'])
                styled_run(_cap, block['title'], size=10.5)
                _cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif t == 'pb':
            br = doc.add_paragraph()
            br.add_run().add_break(WD_BREAK.PAGE)
        elif t == 'title':
            p = new_para(doc, exp['cover'])
            styled_run(p, block['x'], size=block.get('size', 16), bold=True)
        elif t == 'sub':
            p = new_para(doc, exp['cover'])
            styled_run(p, block['x'], size=block.get('size', 10.5), bold=False)
        else:
            raise ValueError('unknown block type: %s' % t)

    doc.save(out_path)
    print('saved: %s' % out_path)
    print('paragraphs=%d tables=%d' % (len(doc.paragraphs), len(doc.tables)))

if __name__ == '__main__':
    main()
