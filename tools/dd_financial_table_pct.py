# -*- coding: utf-8 -*-
"""
财报表格加「占比」列 + 利润表加「毛利额」行（可复用工具，防"加列错位"坑）
用法：
  python dd_financial_table_pct.py <docx> --balance 9,12 --profit 10,13
  （--balance 资产负债表表号；--profit 利润表表号；默认会自动备份）

关键修复（本次复盘沉淀）：
  1. 先快照全部数值再插入，避免"插入后 cell(row,col) 网格错位"读错单元格
  2. pct() 记得 ×100（比例→百分数字符串）
  3. 插入列后 2025年度 从 col2 变 col3，须用新快照的列索引
  4. 克隆原单元格/行保格式（数字右对齐+TNR+%）
口径：资产项÷总资产、负债项÷总负债、权益项÷权益、利润表各项÷营业收入；资产负债率行="—"；毛利额=收入-成本。
"""
import sys, shutil, datetime, argparse
from copy import deepcopy
from docx import Document
from docx.oxml.ns import qn


def celltext(cell): return " ".join(p.text for p in cell.paragraphs if p.text).strip()
def parse(x):
    try: return float(x.replace(",", "").replace("%", "").replace("亿元", ""))
    except Exception: return None
def pct(v): return f"{v*100:.2f}%"   # 修复：×100

def set_tc_text(tc, text):
    ps = tc.findall(qn("w:p"))
    if not ps: return
    p = ps[0]; runs = p.findall(qn("w:r"))
    if not runs: return
    ts = runs[0].findall(qn("w:t"))
    if ts: ts[0].text = text
    for r in runs[1:]:
        for t in r.findall(qn("w:t")): t.text = ""

def snapshot(doc, ti):
    t = doc.tables[ti]
    return [[celltext(t.cell(r, c)) for c in range(len(t.columns))] for r in range(len(t.rows))]

def insert_col_after(doc, ti, col_idx, values):
    t = doc.tables[ti]
    for r in range(len(t.rows)):
        row = t.rows[r]._tr
        tcs = row.findall(qn("w:tc"))
        ref = tcs[col_idx]
        new_tc = deepcopy(ref)
        set_tc_text(new_tc, values.get(r, "—"))
        ref.addnext(new_tc)
    grid = t._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        gcs = grid.findall(qn("w:gridCol"))
        if len(gcs) > col_idx:
            gcs[col_idx].addnext(deepcopy(gcs[col_idx]))

def do_balance(doc, ti):
    data = snapshot(doc, ti)
    def findrow(lab):
        for r in range(len(data)):
            if data[r][0] == lab: return r
        return None
    i_total, i_liab, i_eq = findrow("资产总计"), findrow("负债合计"), findrow("所有者权益合计")
    total, liab, eq = parse(data[i_total][1]), parse(data[i_liab][1]), parse(data[i_eq][1])
    vals = {}
    for r in range(len(data)):
        lab, v = data[r][0], parse(data[r][1])
        if lab in ("资产总计", "负债合计", "所有者权益合计", "负债和所有者权益总计"): val = "100.00%"
        elif lab == "资产负债率" or v is None: val = "—"
        elif r < i_total: val = pct(v / total)
        elif r < i_liab:  val = pct(v / liab)
        elif r < i_eq:    val = pct(v / eq)
        else:             val = "—"
        vals[r] = val
    vals[0] = "占比"
    insert_col_after(doc, ti, 1, vals)
    print(f"  [资产负债表表{ti}] 总资产={total} 负债={liab} 权益={eq} 占比列已加")

def do_profit(doc, ti):
    data = snapshot(doc, ti)
    i_rev = next((r for r in range(len(data)) if data[r][0] == "营业收入"), 0)
    i_cost = next((r for r in range(len(data)) if data[r][0] == "营业成本"), 1)
    rev = parse(data[i_rev][1])
    vals = {}
    for r in range(len(data)):
        lab, v = data[r][0], parse(data[r][1])
        if lab == "营业收入": val = "100.00%"
        elif v is None: val = "—"
        else: val = pct(v / rev)
        vals[r] = val
    vals[0] = "占比"
    insert_col_after(doc, ti, 1, vals)
    # 用新快照(占比列已插入)取 2025年度=col3
    data2 = snapshot(doc, ti)
    i_rev2 = next((r for r in range(len(data2)) if data2[r][0] == "营业收入"), 0)
    i_cost2 = next((r for r in range(len(data2)) if data2[r][0] == "营业成本"), 1)
    rev2, cost2 = parse(data2[i_rev2][1]), parse(data2[i_cost2][1])
    rev25, cost25 = parse(data2[i_rev2][3]), parse(data2[i_cost2][3])   # 列3=2025年度
    gross = rev2 - cost2
    gross25 = (rev25 - cost25) if (rev25 is not None and cost25 is not None) else None
    gross_chg = (gross / (gross25 / 4) - 1) if gross25 else None
    cost_tr = doc.tables[ti].rows[i_cost2]._tr
    new_tr = deepcopy(cost_tr); cost_tr.addnext(new_tr)
    cells = new_tr.findall(qn("w:tc"))
    new_vals = ["毛利额", f"{gross:.2f}", pct(gross / rev2),
                f"{gross25:.2f}" if gross25 is not None else "—",
                (f"+{gross_chg*100:.1f}%" if gross_chg is not None else "—")]
    for i, txt in enumerate(new_vals):
        if i < len(cells): set_tc_text(cells[i], txt)
    print(f"  [利润表表{ti}] 营收={rev} 毛利额2026Q1={gross:.2f} 毛利率={gross/rev2*100:.2f}% 2025={gross25} 变动={gross_chg*100:.1f}%")

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("docx"); ap.add_argument("--balance", default="", help="资产负债表表号,逗号分隔")
    ap.add_argument("--profit", default="", help="利润表表号,逗号分隔")
    a = ap.parse_args()
    if a.balance: a.balance = [int(x) for x in a.balance.split(",") if x.strip()]
    if a.profit:  a.profit  = [int(x) for x in a.profit.split(",") if x.strip()]
    bak = a.docx.replace(".docx", f"_pct_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.docx")
    shutil.copy2(a.docx, bak); print("备份 ->", bak)
    doc = Document(a.docx)
    for ti in (a.balance or []): do_balance(doc, ti)
    for ti in (a.profit or []):  do_profit(doc, ti)
    doc.save(a.docx); print("已保存:", a.docx)

if __name__ == "__main__":
    main()
