# -*- coding: utf-8 -*-
# dd_extract_xlsx.py <input.xlsx> <output.txt>
# 通用xlsx提取：导出全部sheet的非空行（data_only），附合并单元格信息。
import sys
import openpyxl

def fmt(v):
    if v is None:
        return ''
    if isinstance(v, float):
        if v == int(v):
            return str(int(v))
        return ('%.4f' % v).rstrip('0').rstrip('.')
    return str(v).replace('\n', ' / ').strip()

def main():
    src, dst = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(src, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append('##### SHEET: %s (max_row=%d, max_col=%d) #####' % (ws.title, ws.max_row, ws.max_column))
        merged = [str(r) for r in ws.merged_cells.ranges]
        if merged:
            out.append('[merged] ' + ', '.join(merged[:50]))
        for row in ws.iter_rows():
            vals = [fmt(c.value) for c in row]
            while vals and vals[-1] == '':
                vals.pop()
            if vals:
                out.append('R%d: ' % row[0].row + ' | '.join(vals))
        out.append('')
    with open(dst, 'w', encoding='utf-8') as f:
        f.write('\n'.join(out))
    print('xlsx dump done, sheets=%d' % len(wb.worksheets))

if __name__ == '__main__':
    main()
